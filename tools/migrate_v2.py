import sys
import openpyxl, re, json, os, datetime

f1 = r'D:\全景项目\常州中级人民法院\资产\常州中院-云桌面帐号及IP规划.xlsx'
f2 = r'D:\全景项目\常州中级人民法院\资产\中院云桌面资产标签.xlsx'

wb1 = openpyxl.load_workbook(f1, data_only=True)
wb2 = openpyxl.load_workbook(f2, data_only=True)

# 1. 部门映射
user_dept = {}
for sheet in ['专网7段', '专网8段', '专网9段', '专网10段']:
    ws = wb1[sheet]
    for r in list(ws.iter_rows(values_only=True))[1:]:
        if len(r) >= 4 and r[1] and r[3]:
            d = str(r[1]).strip()
            u = str(r[3]).strip()
            if d and u and '未使用' not in u:
                user_dept[u] = d

# 2. 盒子信息映射 (SN -> 终端IP, 状态, 接入时间)
ws_hezi = wb2['盒子登记']
hezi_map = {}
for r in list(ws_hezi.iter_rows(values_only=True))[1:]:
    if r[0]:
        sn = str(r[0]).strip()
        raw_st = str(r[1]).strip() if r[1] else ''
        status = '在用'
        if '闲置' in raw_st: status = '闲置'
        elif '损坏' in raw_st: status = '维修中'
        elif '未拆' in raw_st: status = '仓库备用'
        elif '仓库' in raw_st: status = '仓库备用'
        elif '在用' in raw_st: status = '在用'

        hezi_map[sn] = {
            'status': status,
            'terminal_ip': str(r[5]).strip() if r[5] else '',
            'last_conn': str(r[4]).strip() if r[4] else ''
        }

def fmt_ip(val):
    if not val: return ''
    s = str(val).strip()
    if s.startswith('138.') or s.startswith('192.') or s.startswith('180.') or s.startswith('191.'):
        return s
    if re.match(r'^\d+\.\d+$', s):
        return '138.64.' + s
    return s

# 3. 提取资产主数据
ws_dengji = wb2['登记']
assets = []
asset_logs = []
log_id = 1
asset_id = 1

for r in list(ws_dengji.iter_rows(values_only=True))[1:]:
    if not any(r[:11]):
        continue
    code = str(r[0]).strip() if r[0] else ''
    disp_sn = str(r[1]).strip() if r[1] else ''
    box_sn = str(r[2]).strip() if r[2] else ''
    room = str(r[7]).strip() if r[7] else ''
    sec = str(r[8]).strip() if r[8] else ''
    sec_soft = '已安装' if sec.upper() == 'Y' else ('未安装' if sec.upper() == 'N' else sec)
    ip = fmt_ip(r[9])
    name = str(r[10]).strip() if r[10] else ''
    remark = str(r[11]).strip() if r[11] else ''

    dept = user_dept.get(name, '')
    box_meta = hezi_map.get(box_sn, {})
    term_ip = box_meta.get('terminal_ip', '')
    status = box_meta.get('status', '在用' if name else '闲置')

    asset_item = {
        'id': asset_id,
        'asset_code': code,
        'box_sn': box_sn,
        'display_sn': disp_sn,
        'user_name': name,
        'department': dept,
        'room': room,
        'ip': ip,
        'terminal_ip': term_ip,
        'security_software': sec_soft,
        'status': status,
        'net_type': '专网',
        'remark': remark,
        'created_at': '2026-03-01 09:00:00',
        'updated_at': '2026-03-25 10:00:00'
    }
    assets.append(asset_item)

    desc_user = name if name else '公用/备用'
    asset_logs.append({
        'id': log_id,
        'asset_id': asset_id,
        'log_type': '建档入库',
        'from_user': '',
        'to_user': name,
        'from_location': '',
        'to_location': (dept + ' ' + room).strip(),
        'old_sn': '',
        'new_sn': box_sn,
        'operator': '系统初始化',
        'description': f'资产初次建档登记，分配使用人：{desc_user}',
        'log_time': '2026-03-01 09:00:00'
    })
    log_id += 1
    asset_id += 1

# 4. 维修替换
ws_wx = wb2['维修替换']
for r in list(ws_wx.iter_rows(values_only=True))[1:]:
    if any(r):
        m_type = str(r[0]).strip() if r[0] else '云桌面盒子'
        old_sn = str(r[1]).strip() if r[1] else ''
        new_sn = str(r[2]).strip() if r[2] else ''
        out_time = str(r[3])[:10] if r[3] else '2026-03-31'
        in_time = str(r[4])[:10] if r[4] else '2026-04-03'

        matched_id = None
        for a in assets:
            if a['box_sn'] == old_sn or a['box_sn'] == new_sn:
                matched_id = a['id']
                break

        asset_logs.append({
            'id': log_id,
            'asset_id': matched_id if matched_id else 0,
            'log_type': '设备维修更换',
            'from_user': '',
            'to_user': '',
            'from_location': '',
            'to_location': '',
            'old_sn': old_sn,
            'new_sn': new_sn,
            'operator': '运维技术员',
            'description': f'{m_type}换件维修（原SN: {old_sn} 替换为 新SN: {new_sn}，出库: {out_time}，入库完成: {in_time}）',
            'log_time': f'{out_time} 14:30:00'
        })
        log_id += 1

# 5. 回收记录
ws_hs = wb2['云桌面回收']
for r in list(ws_hs.iter_rows(values_only=True))[1:]:
    if any(r):
        u_name = str(r[0]).strip() if r[0] else ''
        u_ip = str(r[1]).strip() if r[1] else ''
        c_time = str(r[2])[:10] if r[2] else '2026-03-31'
        is_occ = str(r[3]).strip() if r[3] else ''

        matched_id = None
        for a in assets:
            if a['user_name'] == u_name or a['ip'] == u_ip:
                matched_id = a['id']
                break

        asset_logs.append({
            'id': log_id,
            'asset_id': matched_id if matched_id else 0,
            'log_type': '设备回收注销',
            'from_user': u_name,
            'to_user': '仓库',
            'from_location': '',
            'to_location': '资产库房',
            'old_sn': '',
            'new_sn': '',
            'operator': '资产管理员',
            'description': f'云桌面回收下线（原使用人: {u_name}，IP: {u_ip}，是否占用: {is_occ}）',
            'log_time': f'{c_time} 10:00:00'
        })
        log_id += 1

# 6. IP池
ip_pool = []
ip_id = 1
for sheet in ['专网7段', '专网8段', '专网9段', '专网10段']:
    ws = wb1[sheet]
    subnet = sheet.replace('专网', '')
    for r in list(ws.iter_rows(values_only=True))[1:]:
        if len(r) >= 4 and r[2]:
            dept = str(r[1]).strip() if r[1] else ''
            ip = str(r[2]).strip()
            acct = str(r[3]).strip() if r[3] else ''
            rem = str(r[4]).strip() if len(r) > 4 and r[4] else ''
            status = '空闲' if ('未使用' in acct or acct == '') else '在用'
            ip_pool.append({
                'id': ip_id,
                'subnet': subnet,
                'ip': ip,
                'department': dept,
                'room': '',
                'account': acct,
                'status': status,
                'remark': rem
            })
            ip_id += 1

# 外网IP
ws_ww = wb1['外网 IP']
for r in list(ws_ww.iter_rows(values_only=True))[2:]:
    if len(r) >= 6 and (r[1] or r[2]):
        box_ip = str(r[1]).strip() if r[1] else ''
        biz_ip = str(r[2]).strip() if r[2] else ''
        dept = str(r[3]).strip() if r[3] else ''
        room = str(r[4]).strip() if r[4] else ''
        user = str(r[5]).strip() if r[5] else ''
        status = '在用' if user else '空闲'
        ip_pool.append({
            'id': ip_id,
            'subnet': '外网',
            'ip': f'{biz_ip} (盒子:{box_ip})' if box_ip else biz_ip,
            'department': dept,
            'room': room,
            'account': user,
            'status': status,
            'remark': '外网云桌面'
        })
        ip_id += 1

out_data = {
    'version': 2,
    'assets': assets,
    'asset_logs': asset_logs,
    'ip_pool': ip_pool
}

seed_path = r'D:\全景项目\常州中级人民法院\资产\seed\seed_v2.json'
with open(seed_path, 'w', encoding='utf-8') as f:
    json.dump(out_data, f, ensure_ascii=False, indent=2)

print(f'Done! Assets: {len(assets)}, Logs: {len(asset_logs)}, IP pool: {len(ip_pool)}')
